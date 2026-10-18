"""Тесты ЛР5.

Запуск из корня репозитория:
    python3 -m unittest discover -s tests -v
"""
import csv
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from src.lab05.csv_xlsx import csv_to_xlsx
from src.lab05.json_csv import csv_to_json, json_to_csv


class Base(unittest.TestCase):
    """Базовый класс: создаёт временную папку для файлов каждого теста."""

    def setUp(self):
        """Создаёт временную папку перед тестом."""
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)

    def tearDown(self):
        """Удаляет временную папку после теста."""
        self._tmp.cleanup()

    def write(self, name, text):
        """Создаёт во временной папке файл с текстом.

        Args:
            name (str): Имя файла.
            text (str): Содержимое файла.

        Returns:
            str: Путь к созданному файлу.
        """
        p = self.dir / name
        p.write_text(text, encoding="utf-8")
        return str(p)

    def p(self, name):
        """Возвращает путь к файлу во временной папке (файл может не существовать).

        Args:
            name (str): Имя файла.

        Returns:
            str: Путь к файлу.
        """
        return str(self.dir / name)


class TestJsonToCsv(Base):
    """Тесты функции json_to_csv."""

    def test_basic(self):
        """Обычный JSON со списком словарей превращается в CSV, кириллица сохраняется."""
        src = self.write("a.json", json.dumps([{"name": "Иван", "age": 30}, {"name": "Bob", "age": 25}], ensure_ascii=False))
        json_to_csv(src, self.p("a.csv"))
        with open(self.p("a.csv"), newline="", encoding="utf-8") as f:
            rows = list(csv.DictReader(f))
        self.assertEqual(rows, [{"name": "Иван", "age": "30"}, {"name": "Bob", "age": "25"}])

    def test_missing_fields_filled_with_empty(self):
        """Если у объекта нет какого-то поля, в CSV на его месте пустая строка."""
        src = self.write("a.json", '[{"a": 1, "b": 2}, {"a": 3, "c": 4}]')
        json_to_csv(src, self.p("a.csv"))
        with open(self.p("a.csv"), newline="", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            rows = list(reader)
        self.assertEqual(reader.fieldnames, ["a", "b", "c"])
        self.assertEqual(rows[0]["c"], "")
        self.assertEqual(rows[1]["b"], "")

    def test_empty_list(self):
        """Пустой список [] в JSON вызывает ValueError."""
        with self.assertRaises(ValueError):
            json_to_csv(self.write("a.json", "[]"), self.p("a.csv"))

    def test_not_a_list(self):
        """JSON, где на верхнем уровне не список (а объект), вызывает ValueError."""
        with self.assertRaises(ValueError):
            json_to_csv(self.write("a.json", '{"a": 1}'), self.p("a.csv"))

    def test_list_of_non_dicts(self):
        """Список из чисел вместо словарей вызывает ValueError."""
        with self.assertRaises(ValueError):
            json_to_csv(self.write("a.json", "[1, 2, 3]"), self.p("a.csv"))

    def test_broken_json(self):
        """Файл с синтаксической ошибкой в JSON вызывает ValueError."""
        with self.assertRaises(ValueError):
            json_to_csv(self.write("a.json", "{oops"), self.p("a.csv"))

    def test_wrong_extension(self):
        """Входной файл не .json вызывает ValueError."""
        with self.assertRaises(ValueError):
            json_to_csv(self.write("a.txt", "[]"), self.p("a.csv"))

    def test_file_not_found(self):
        """Несуществующий входной файл вызывает FileNotFoundError."""
        with self.assertRaises(FileNotFoundError):
            json_to_csv(self.p("nope.json"), self.p("a.csv"))


class TestCsvToJson(Base):
    """Тесты функции csv_to_json."""

    def test_basic_values_are_strings(self):
        """Обычный CSV превращается в список словарей, все значения остаются строками."""
        src = self.write("a.csv", "name,age\nAlice,22\nМария,28\n")
        csv_to_json(src, self.p("a.json"))
        data = json.loads(Path(self.p("a.json")).read_text(encoding="utf-8"))
        self.assertEqual(data, [{"name": "Alice", "age": "22"}, {"name": "Мария", "age": "28"}])

    def test_semicolon_delimiter(self):
        """CSV с разделителем «;» читается так же, как с запятой."""
        src = self.write("a.csv", "name;age\nAlice;22\nBob;25\n")
        csv_to_json(src, self.p("a.json"))
        data = json.loads(Path(self.p("a.json")).read_text(encoding="utf-8"))
        self.assertEqual(data[1], {"name": "Bob", "age": "25"})

    def test_ensure_ascii_false(self):
        """В JSON кириллица записана буквами, а не кодами \\uXXXX."""
        src = self.write("a.csv", "name\nИван\n")
        csv_to_json(src, self.p("a.json"))
        self.assertIn("Иван", Path(self.p("a.json")).read_text(encoding="utf-8"))

    def test_empty_file(self):
        """Пустой CSV-файл вызывает ValueError."""
        with self.assertRaises(ValueError):
            csv_to_json(self.write("a.csv", ""), self.p("a.json"))

    def test_header_only(self):
        """CSV только с заголовком, без строк данных, вызывает ValueError."""
        with self.assertRaises(ValueError):
            csv_to_json(self.write("a.csv", "name,age\n"), self.p("a.json"))

    def test_wrong_extension(self):
        """Входной файл не .csv вызывает ValueError."""
        with self.assertRaises(ValueError):
            csv_to_json(self.write("a.txt", "a\n1\n"), self.p("a.json"))

    def test_file_not_found(self):
        """Несуществующий входной файл вызывает FileNotFoundError."""
        with self.assertRaises(FileNotFoundError):
            csv_to_json(self.p("nope.csv"), self.p("a.json"))

    def test_round_trip_count(self):
        """CSV -> JSON -> CSV: количество записей до и после совпадает (3 и 3)."""
        src = self.write("a.csv", "x,y\n1,2\n3,4\n5,6\n")
        csv_to_json(src, self.p("a.json"))
        json_to_csv(self.p("a.json"), self.p("b.csv"))
        with open(self.p("b.csv"), newline="", encoding="utf-8") as f:
            self.assertEqual(len(list(csv.DictReader(f))), 3)


class TestCsvToXlsx(Base):
    """Тесты функции csv_to_xlsx."""

    def test_basic(self):
        """CSV превращается в XLSX: лист Sheet1, верный заголовок, все строки на месте."""
        src = self.write("a.csv", "name,city\nAlice,SPB\nИван,Екатеринбург\n")
        csv_to_xlsx(src, self.p("a.xlsx"))
        ws = load_workbook(self.p("a.xlsx")).active
        self.assertEqual(ws.title, "Sheet1")
        self.assertEqual([c.value for c in ws[1]], ["name", "city"])
        self.assertEqual(ws.max_row, 3)
        self.assertEqual(ws["B3"].value, "Екатеринбург")

    def test_min_width_and_autowidth(self):
        """Ширина колонки не меньше 8, а длинный заголовок делает колонку шире короткой."""
        src = self.write("a.csv", "a,long_header_name\n1,2\n")
        csv_to_xlsx(src, self.p("a.xlsx"))
        ws = load_workbook(self.p("a.xlsx")).active
        self.assertGreaterEqual(ws.column_dimensions["A"].width, 8)
        self.assertGreater(ws.column_dimensions["B"].width, ws.column_dimensions["A"].width)

    def test_empty_csv(self):
        """Пустой CSV-файл вызывает ValueError."""
        with self.assertRaises(ValueError):
            csv_to_xlsx(self.write("a.csv", ""), self.p("a.xlsx"))

    def test_wrong_extension(self):
        """Неверное расширение входного (не .csv) или выходного (не .xlsx) файла вызывает ValueError."""
        with self.assertRaises(ValueError):
            csv_to_xlsx(self.write("a.txt", "a\n1\n"), self.p("a.xlsx"))
        with self.assertRaises(ValueError):
            csv_to_xlsx(self.write("a.csv", "a\n1\n"), self.p("a.xls"))

    def test_file_not_found(self):
        """Несуществующий входной файл вызывает FileNotFoundError."""
        with self.assertRaises(FileNotFoundError):
            csv_to_xlsx(self.p("nope.csv"), self.p("a.xlsx"))


if __name__ == "__main__":
    unittest.main()