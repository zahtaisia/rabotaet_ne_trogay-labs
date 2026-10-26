"""Конвертация CSV -> XLSX через openpyxl (ЛР5, задание B)."""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
import sys

MIN_WIDTH = 8


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """Конвертирует CSV-файл в XLSX.

    Первая строка CSV — заголовок (выделяется жирным). Лист называется
    "Sheet1". Ширина колонки равна длине самого длинного значения
    (плюс небольшой запас), но не менее 8 символов. Разделитель
    определяется автоматически.

    Args:
        csv_path (str): Путь к входному .csv файлу.
        xlsx_path (str): Путь к выходному .xlsx файлу.

    Raises:
        ValueError: Если расширение входного или выходного файла неверное,
            либо CSV пустой или без заголовка.
        FileNotFoundError: Если входной файл не существует.
    """
    src = Path(csv_path)
    dst = Path(xlsx_path)

    if src.suffix.lower() != ".csv":
        raise ValueError(f"Неверный тип входного файла: '{src.suffix}', нужен .csv")
    if dst.suffix.lower() != ".xlsx":
        raise ValueError(f"Неверный тип выходного файла: '{dst.suffix}', нужен .xlsx")
    if not src.is_file():
        raise FileNotFoundError(f"Файл не найден: {src}")

    with src.open(newline="", encoding="utf-8-sig") as f:
        sample = f.read(4096)
        f.seek(0)
        if not sample.strip():
            raise ValueError("Пустой CSV")
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = [row for row in csv.reader(f, dialect=dialect)]

    if not rows or not any(cell.strip() for cell in rows[0]):
        raise ValueError("CSV без заголовка или пустой")

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    for row in rows:
        ws.append(row)

    for cell in ws[1]:  # выделяем заголовок
        cell.font = Font(bold=True)

    col_count = max(len(r) for r in rows)
    for idx in range(1, col_count + 1):
        longest = max(
            (len(r[idx - 1]) for r in rows if idx - 1 < len(r)), default=0
        )
        ws.column_dimensions[get_column_letter(idx)].width = max(MIN_WIDTH, longest + 2)

    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    
    if __name__ == "__main__":
        try:
            csv_to_xlsx("data/lab05/samples/people.csv", "data/lab05/out/people.xlsx")
            csv_to_xlsx("data/lab05/samples/cities.csv", "data/lab05/out/cities.xlsx")
        except (FileNotFoundError, ValueError) as e:
            print(f"Ошибка: {e}", file=sys.stderr)
            sys.exit(1)