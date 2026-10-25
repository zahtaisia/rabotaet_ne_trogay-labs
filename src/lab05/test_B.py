from src.lab05.csv_xlsx import csv_to_xlsx
from openpyxl import load_workbook

csv_to_xlsx("data/lab05/samples/people.csv", "data/lab05/out/people.xlsx")
csv_to_xlsx("data/lab05/samples/cities.csv", "data/lab05/out/cities.xlsx")

wb = load_workbook("data/lab05/out/people.xlsx")
ws = wb.active
print("Лист:", ws.title)
print("XLSX строк с данными:", ws.max_row - 1)
print("Колонок:", ws.max_column)
for letter, dim in ws.column_dimensions.items():
    print("Ширина колонки", letter, "=", dim.width)